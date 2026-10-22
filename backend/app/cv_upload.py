"""Bounded text extraction for synthetic PDF and DOCX resumes.

This parser does not persist the supplied bytes or extracted text and never
uses the client filename as a filesystem path. Multipart handling may spool
the upload to a temporary file before this parser runs.
"""
from __future__ import annotations

from io import BytesIO
import os
import time
from pathlib import PurePath
from zipfile import ZipFile

import fitz
from docx import Document

from .services import APIError

MAX_FILE_BYTES = 10 * 1024 * 1024
VERCEL_FILE_BYTES = 3 * 1024 * 1024
MAX_TEXT_CHARS = 50_000
MAX_PDF_PAGES = 30
MAX_PARSE_SECONDS = 10
MAX_ZIP_ENTRIES = 500
MAX_UNCOMPRESSED_BYTES = 30 * 1024 * 1024
MAX_ENTRY_BYTES = 10 * 1024 * 1024


def upload_limit() -> int:
    maximum = VERCEL_FILE_BYTES if os.getenv('VERCEL') else MAX_FILE_BYTES
    try:
        value = int(os.getenv('MAX_CV_UPLOAD_BYTES', str(maximum)))
    except ValueError:
        value = maximum
    return max(1, min(value, maximum))


def _bad(message: str, code: str = 'CV_INVALID_FILE', status: int = 422) -> APIError:
    return APIError(status, code, message)


def _pdf_text(data: bytes) -> str:
    if not data.startswith(b'%PDF-'):
        raise _bad('File content is not a PDF')
    started = time.monotonic()
    try:
        with fitz.open(stream=data, filetype='pdf') as doc:
            if doc.needs_pass:
                raise _bad('Encrypted PDFs are not supported')
            if doc.page_count > MAX_PDF_PAGES:
                raise _bad('PDF has too many pages', 'CV_TOO_COMPLEX')
            chunks = []
            total = 0
            for page in doc:
                if time.monotonic() - started > MAX_PARSE_SECONDS:
                    raise _bad('PDF parsing took too long', 'CV_TOO_COMPLEX')
                text = page.get_text('text')
                if time.monotonic() - started > MAX_PARSE_SECONDS:
                    raise _bad('PDF parsing took too long', 'CV_TOO_COMPLEX')
                total += len(text)
                if total > MAX_TEXT_CHARS:
                    raise _bad('Extracted text is too long', 'CV_TOO_COMPLEX')
                chunks.append(text)
            result = '\n'.join(chunks).strip()
    except APIError:
        raise
    except Exception:
        raise _bad('PDF is damaged or cannot be parsed') from None
    if not result:
        raise _bad('No extractable text found; upload a text-based PDF instead of a scan', 'CV_NO_TEXT')
    return result


def _docx_text(data: bytes) -> str:
    if not data.startswith(b'PK\x03\x04'):
        raise _bad('File content is not a DOCX')
    started = time.monotonic()
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise _bad('DOCX has too many parts', 'CV_TOO_COMPLEX')
            names = {item.filename for item in entries}
            if '[Content_Types].xml' not in names or 'word/document.xml' not in names:
                raise _bad('DOCX package is invalid')
            total = 0
            for item in entries:
                total += item.file_size
                if (item.file_size > MAX_ENTRY_BYTES or total > MAX_UNCOMPRESSED_BYTES
                    or item.flag_bits & 1 or item.filename.startswith('/')
                    or '..' in PurePath(item.filename).parts
                    or item.filename.lower().endswith(('.bin', '.exe', '.js'))
                    or '/embeddings/' in item.filename.lower()):
                    raise _bad('DOCX contains unsupported or oversized parts', 'CV_TOO_COMPLEX')
            if archive.testzip() is not None:
                raise _bad('DOCX package is damaged')
        document = Document(BytesIO(data))
        chunks = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                chunks.extend(cell.text for cell in row.cells)
        result = '\n'.join(chunks).strip()
        if time.monotonic() - started > MAX_PARSE_SECONDS:
            raise _bad('DOCX parsing took too long', 'CV_TOO_COMPLEX')
    except APIError:
        raise
    except Exception:
        raise _bad('DOCX is damaged or cannot be parsed') from None
    if len(result) > MAX_TEXT_CHARS:
        raise _bad('Extracted text is too long', 'CV_TOO_COMPLEX')
    if not result:
        raise _bad('No extractable text found in DOCX', 'CV_NO_TEXT')
    return result


def extract_cv_text(filename: str | None, data: bytes) -> str:
    # Only the extension is used; never open or store the supplied filename.
    suffix = PurePath(filename or '').suffix.lower()
    if suffix not in ('.pdf', '.docx'):
        raise _bad('Only PDF and DOCX files are supported', 'CV_UNSUPPORTED_TYPE', 415)
    if not data:
        raise _bad('Uploaded file is empty')
    if len(data) > upload_limit():
        raise _bad('CV exceeds the upload size limit', 'CV_TOO_LARGE', 413)
    return _pdf_text(data) if suffix == '.pdf' else _docx_text(data)
